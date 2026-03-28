from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse, Response
import os
import httpx
from sqlalchemy.orm import Session
from sqlalchemy import func
from pydantic import BaseModel
from database import get_db
from models import Lead, EmailLog
import io

router = APIRouter()


# ── Pydantic schemas ──────────────────────────────────────────────────────────

class LeadCreate(BaseModel):
    business_name: str
    category: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    website: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    country: Optional[str] = "US"
    lat: Optional[float] = None
    lng: Optional[float] = None
    contact_name: Optional[str] = None
    position: Optional[str] = None
    notes: Optional[str] = None
    source: Optional[str] = "manual"


class LeadUpdate(BaseModel):
    business_name: Optional[str] = None
    category: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    website: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    state: Optional[str] = None
    country: Optional[str] = None
    lat: Optional[float] = None
    lng: Optional[float] = None
    contact_name: Optional[str] = None
    position: Optional[str] = None
    status: Optional[str] = None
    notes: Optional[str] = None
    archived: Optional[bool] = None


class BulkFindEmailRequest(BaseModel):
    lead_ids: list[int]


class StatusUpdate(BaseModel):
    status: str


# ── Routes ────────────────────────────────────────────────────────────────────

@router.get("/")
def list_leads(
    db: Session = Depends(get_db),
    country: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    city: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    has_email: Optional[bool] = Query(None),
    email_source: Optional[str] = Query(None),   # "hunter" | "scraper" | "none"
    sort: Optional[str] = Query(None),            # "score_desc" | "score_asc" | "name_asc"
    archived: Optional[bool] = Query(None),
    skip: int = 0,
    limit: int = 500,
):
    q = db.query(Lead)
    if archived is True:
        q = q.filter(Lead.archived == True)   # noqa: E712
    else:
        q = q.filter(Lead.archived == False)  # noqa: E712
    if country:   q = q.filter(Lead.country == country)
    if state:     q = q.filter(Lead.state == state)
    if city:      q = q.filter(Lead.city.ilike(f"%{city}%"))
    if category:  q = q.filter(Lead.category.ilike(f"%{category}%"))
    if status:    q = q.filter(Lead.status == status)
    if has_email is True:  q = q.filter(Lead.email.isnot(None), Lead.email != "")
    if has_email is False: q = q.filter((Lead.email.is_(None)) | (Lead.email == ""))
    if email_source == "hunter":  q = q.filter(Lead.email_source == "hunter")
    if email_source == "scraper": q = q.filter(Lead.email_source == "scraper")
    if email_source == "none":    q = q.filter((Lead.email.is_(None)) | (Lead.email == ""))
    if search:
        q = q.filter(
            Lead.business_name.ilike(f"%{search}%") |
            Lead.city.ilike(f"%{search}%") |
            Lead.email.ilike(f"%{search}%")
        )
    total = q.count()
    if sort == "name_asc":
        leads = q.order_by(Lead.business_name.asc()).offset(skip).limit(limit).all()
    else:
        leads = q.order_by(Lead.created_at.desc()).offset(skip).limit(limit).all()
    result = [_serialize(l) for l in leads]
    # Score sorting is done in Python since score is computed, not stored
    if sort == "score_desc": result.sort(key=lambda l: l["score"], reverse=True)
    if sort == "score_asc":  result.sort(key=lambda l: l["score"])
    return {"total": total, "leads": result}


@router.get("/map-pins")
def map_pins(db: Session = Depends(get_db)):
    leads = db.query(
        Lead.id, Lead.business_name, Lead.city, Lead.state,
        Lead.lat, Lead.lng, Lead.status, Lead.category
    ).filter(Lead.lat.isnot(None), Lead.lng.isnot(None), Lead.archived == False).all()
    return [
        {
            "id": l.id, "business_name": l.business_name,
            "city": l.city, "state": l.state,
            "lat": l.lat, "lng": l.lng,
            "status": l.status, "category": l.category,
        }
        for l in leads
    ]


@router.get("/stats")
def stats(db: Session = Depends(get_db)):
    total = db.query(func.count(Lead.id)).filter(Lead.archived == False).scalar()
    by_status = db.query(Lead.status, func.count(Lead.id)).filter(Lead.archived == False).group_by(Lead.status).all()
    by_state  = db.query(Lead.state, func.count(Lead.id)).filter(Lead.archived == False).group_by(Lead.state).order_by(func.count(Lead.id).desc()).limit(10).all()
    with_email = db.query(func.count(Lead.id)).filter(Lead.archived == False).filter(Lead.email.isnot(None), Lead.email != "").scalar()
    return {
        "total": total,
        "with_email": with_email,
        "by_status": {s: c for s, c in by_status},
        "top_states": [{"state": s, "count": c} for s, c in by_state],
    }


@router.get("/filters")
def filter_options(
    db: Session = Depends(get_db),
    country: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
):
    q_country = db.query(Lead.country).distinct().filter(Lead.country.isnot(None))
    countries = sorted([r[0] for r in q_country.all() if r[0]])

    q_state = db.query(Lead.state).distinct().filter(Lead.state.isnot(None))
    if country:
        q_state = q_state.filter(Lead.country == country)
    states = sorted([r[0] for r in q_state.all() if r[0]])

    q_city = db.query(Lead.city).distinct().filter(Lead.city.isnot(None))
    if state:
        q_city = q_city.filter(Lead.state == state)
    elif country:
        q_city = q_city.filter(Lead.country == country)
    cities = sorted([r[0] for r in q_city.all() if r[0]])

    categories = db.query(Lead.category).distinct().filter(Lead.category.isnot(None)).all()
    return {
        "countries": countries,
        "states": states,
        "cities": cities,
        "categories": sorted(list(set(
            cat.strip()
            for row in categories if row[0]
            for cat in row[0].split(",")
            if cat.strip()
        ))),
    }


@router.get("/export")
def export_leads(
    db: Session = Depends(get_db),
    country: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    city: Optional[str] = Query(None),
    category: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    lead_id: Optional[int] = Query(None),
):
    """Export leads to Excel with optional filters."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    q = db.query(Lead)
    if lead_id:   q = q.filter(Lead.id == lead_id)
    if country:   q = q.filter(Lead.country == country)
    if state:     q = q.filter(Lead.state == state)
    if city:      q = q.filter(Lead.city.ilike(f"%{city}%"))
    if category:  q = q.filter(Lead.category.ilike(f"%{category}%"))
    if status:    q = q.filter(Lead.status == status)
    leads = q.order_by(Lead.state, Lead.city, Lead.business_name).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "GAVA Leads"

    # Header row styling
    header_fill = PatternFill("solid", fgColor="1A1C2E")
    header_font = Font(bold=True, color="C8A96E", size=11)
    headers = [
        "Business Name", "Category", "Contact Name", "Email", "Email Source",
        "Phone", "Website", "Address", "City", "State", "Country",
        "Status", "Score", "Notes", "Added",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Column widths
    widths = [35, 20, 22, 30, 12, 16, 35, 35, 18, 8, 8, 12, 6, 30, 18]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    # Data rows
    for row, lead in enumerate(leads, 2):
        ws.cell(row=row, column=1,  value=lead.business_name)
        ws.cell(row=row, column=2,  value=lead.category)
        ws.cell(row=row, column=3,  value=lead.contact_name)
        ws.cell(row=row, column=4,  value=lead.email)
        ws.cell(row=row, column=5,  value=lead.email_source or "—")
        ws.cell(row=row, column=6,  value=lead.phone)
        ws.cell(row=row, column=7,  value=lead.website)
        ws.cell(row=row, column=8,  value=lead.address)
        ws.cell(row=row, column=9,  value=lead.city)
        ws.cell(row=row, column=10, value=lead.state)
        ws.cell(row=row, column=11, value=lead.country)
        ws.cell(row=row, column=12, value=lead.status)
        ws.cell(row=row, column=13, value=_score_lead(lead))
        ws.cell(row=row, column=14, value=lead.notes)
        ws.cell(row=row, column=15, value=lead.created_at.strftime("%Y-%m-%d") if lead.created_at else "")
        # Alternate row fill
        if row % 2 == 0:
            fill = PatternFill("solid", fgColor="252844")
            for col in range(1, 16):
                ws.cell(row=row, column=col).fill = fill

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "gava_leads"
    if state:     filename += f"_{state}"
    if city:      filename += f"_{city.replace(' ','_')}"
    if lead_id:   filename += f"_lead{lead_id}"
    filename += ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.post("/{lead_id}/find-email")
async def find_email_hunter(lead_id: int, db: Session = Depends(get_db)):
    """
    Call Hunter.io on demand for a specific lead.
    Only consumes a credit when the user explicitly requests it.
    """
    import os
    from services.hunter_service import find_hr_contact as hunter_find
    from services.web_scraper import find_hr_email as scrape_find

    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    if not lead.website:
        raise HTTPException(status_code=400, detail="This lead has no website — can't search for email")

    if os.getenv("HUNTER_API_KEY"):
        result = await hunter_find(lead.website)
        if result.get("email"):
            lead.email         = result["email"]
            lead.contact_name  = result.get("contact_name") or lead.contact_name
            lead.email_source  = "hunter"
            db.commit()
            return {"email": lead.email, "contact_name": lead.contact_name, "source": "hunter"}

    # Fallback to free scraper
    email = await scrape_find(lead.website)
    if email:
        lead.email        = email
        lead.email_source = "scraper"
        db.commit()
        return {"email": email, "source": "scraper"}

    raise HTTPException(status_code=404, detail="No email found for this business")


@router.post("/bulk-find-email")
async def bulk_find_email(payload: BulkFindEmailRequest, db: Session = Depends(get_db)):
    """
    Run Hunter.io (or free scraper fallback) on multiple leads at once.
    Only processes leads that have a website. Returns per-lead results.
    """
    import os
    from services.hunter_service import find_hr_contact as hunter_find
    from services.web_scraper import find_hr_email as scrape_find

    results = []
    for lead_id in payload.lead_ids:
        lead = db.query(Lead).filter(Lead.id == lead_id).first()
        if not lead or not lead.website:
            results.append({"id": lead_id, "status": "skipped", "reason": "no website"})
            continue
        try:
            if os.getenv("HUNTER_API_KEY"):
                result = await hunter_find(lead.website)
                if result.get("email"):
                    lead.email        = result["email"]
                    lead.contact_name = result.get("contact_name") or lead.contact_name
                    lead.email_source = "hunter"
                    db.commit()
                    results.append({"id": lead_id, "status": "found", "email": lead.email, "source": "hunter"})
                    continue

            email = await scrape_find(lead.website)
            if email:
                lead.email        = email
                lead.email_source = "scraper"
                db.commit()
                results.append({"id": lead_id, "status": "found", "email": email, "source": "scraper"})
            else:
                results.append({"id": lead_id, "status": "not_found"})
        except Exception as e:
            results.append({"id": lead_id, "status": "error", "reason": str(e)})

    found = sum(1 for r in results if r["status"] == "found")
    return {"processed": len(results), "found": found, "results": results}


@router.get("/{lead_id}/photo")
async def get_lead_photo(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead or not lead.photo_ref:
        raise HTTPException(status_code=404, detail="No photo available")
    api_key = os.getenv("GOOGLE_PLACES_API_KEY", "")
    url = f"https://places.googleapis.com/v1/{lead.photo_ref}/media?maxHeightPx=600&maxWidthPx=900&key={api_key}"
    async with httpx.AsyncClient(timeout=10, follow_redirects=True) as client:
        resp = await client.get(url)
    if resp.status_code != 200:
        raise HTTPException(status_code=404, detail="Photo not found")
    return Response(content=resp.content, media_type=resp.headers.get("content-type", "image/jpeg"))


@router.patch("/{lead_id}/status")
def update_status(lead_id: int, payload: StatusUpdate, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    lead.status = payload.status
    db.commit()
    return _serialize(lead)


@router.patch("/{lead_id}/archive")
def archive_lead(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    lead.archived = True
    db.commit()
    return {"id": lead_id, "archived": True}


@router.patch("/{lead_id}/unarchive")
def unarchive_lead(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    lead.archived = False
    db.commit()
    return {"id": lead_id, "archived": False}


@router.get("/{lead_id}")
def get_lead(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    data = _serialize(lead)
    data["email_history"] = [
        {"id": e.id, "subject": e.subject, "body": e.body, "sent_at": e.sent_at.isoformat()}
        for e in lead.emails
    ]
    return data


@router.post("/")
def create_lead(payload: LeadCreate, db: Session = Depends(get_db)):
    lead = Lead(**payload.model_dump())
    db.add(lead)
    db.commit()
    db.refresh(lead)
    return _serialize(lead)


@router.patch("/{lead_id}")
def update_lead(lead_id: int, payload: LeadUpdate, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(lead, field, value)
    db.commit()
    db.refresh(lead)
    return _serialize(lead)


@router.delete("/{lead_id}")
def delete_lead(lead_id: int, db: Session = Depends(get_db)):
    lead = db.query(Lead).filter(Lead.id == lead_id).first()
    if not lead:
        raise HTTPException(status_code=404, detail="Lead not found")
    lead.archived = True
    db.commit()
    return {"archived": True, "id": lead_id}


# ── Helpers ───────────────────────────────────────────────────────────────────

def _score_lead(lead: Lead) -> int:
    """
    Lead quality score 0–10.
    Higher = more contactable and more likely to convert.
    """
    score = 0
    if lead.email:        score += 3
    if lead.phone:        score += 2
    if lead.website:      score += 2
    if lead.contact_name: score += 2
    if lead.email_source == "hunter": score += 1
    return score


def _serialize(lead: Lead) -> dict:
    return {
        "id":            lead.id,
        "business_name": lead.business_name,
        "category":      lead.category,
        "email":         lead.email,
        "email_source":  lead.email_source,
        "phone":         lead.phone,
        "website":       lead.website,
        "address":       lead.address,
        "city":          lead.city,
        "state":         lead.state,
        "country":       lead.country,
        "lat":           lead.lat,
        "lng":           lead.lng,
        "contact_name":  lead.contact_name,
        "position":      lead.position,
        "archived":      bool(lead.archived),
        "status":        lead.status,
        "source":        lead.source,
        "notes":         lead.notes,
        "score":         _score_lead(lead),
        "photo_ref":     lead.photo_ref,
        "created_at":    lead.created_at.isoformat() if lead.created_at else None,
        "last_contacted": lead.last_contacted.isoformat() if lead.last_contacted else None,
    }
